"""Deterministic import of the private crypto-history workbook.

The workbook is evidence, not an order feed.  Every source row becomes an
immutable dated analysis.  Exact wording, units, operators and source cells
are retained; ambiguous text is reported instead of guessed.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import asdict, dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from .store import (
    LexaActionRow,
    LexaAnalysisRow,
    LexaConditionRow,
    LexaImportRow,
    LexaLevelRow,
    LexaVideoRow,
    lexa_session,
)

IMPORT_VERSION = "excel-history/1"
REQUIRED_SHEETS = ("Dashboard", "Ordres_actifs", "Historique", "Matrice_dates")
EMPTY = {"", "—", "-", "–"}


@dataclass(slots=True)
class PriceSpec:
    value: float
    value_high: float | None = None
    unit: str = "USD"
    operator: str = "AT"
    approximate: bool = False
    raw: str = ""


@dataclass(slots=True)
class ActionDraft:
    action: str
    amount_type: str = "NONE"
    amount: float | None = None
    raw_text: str = ""
    condition_text: str = ""
    execution_enabled: bool = True
    extra: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class ConditionDraft:
    condition_type: str
    timeframe: str | None
    operator: str
    description: str


@dataclass(slots=True)
class LevelDraft:
    kind: str
    value: float
    value_high: float | None
    unit: str
    label: str
    source_text: str
    confidence: str = "HIGH"
    conditions: list[ConditionDraft] = field(default_factory=list)
    action: ActionDraft | None = None
    extra: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class AnalysisDraft:
    asset: str
    published_at: datetime
    title: str
    source_ref: str
    source_type: str
    source_status: str
    price_at_video: float | None
    stance: str
    summary: str
    market_context: str
    requires_revalidation: bool
    levels: list[LevelDraft]
    extra: dict[str, Any]
    completed: bool = False


@dataclass(slots=True)
class WorkbookDraft:
    path: Path
    fingerprint: str
    analyses: list[AnalysisDraft]
    historical_count: int
    current_count: int
    sheets: dict[str, dict[str, Any]]
    assets: list[str]
    dates: list[str]
    ambiguities: list[dict[str, str]]
    checks: list[dict[str, Any]]


@dataclass(slots=True)
class AmountSpec:
    amount_type: str
    amount: float | None
    raw: str


_PRICE = re.compile(
    r"(?<![\d/])(?P<op>[>~])?\s*"
    r"(?P<a>\d+(?:[ \u202f]\d{3})*(?:[.,]\d+)?)\s*(?P<ak>[kK])?"
    r"(?:\s*[–-]\s*(?P<b>\d+(?:[ \u202f]\d{3})*(?:[.,]\d+)?)\s*(?P<bk>[kK])?)?"
    r"\s*(?P<unit>[$€])?"
)
_AMOUNT = re.compile(r"(?P<n>\d+(?:[ \u202f]\d{3})*(?:[.,]\d+)?)\s*(?P<u>€|\$|%)")


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _date(value: Any, *, cell: str, ambiguities: list[dict[str, str]]) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=UTC)
    raw = _text(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt).replace(tzinfo=UTC)
        except ValueError:
            pass
    ambiguities.append({"cell": cell, "value": raw, "reason": "Date illisible"})
    return None


def _number(raw: str, reference: float | None = None, kilo: bool = False) -> float:
    clean = raw.replace(" ", "").replace("\u202f", "")
    if "," in clean:
        left, right = clean.split(",", 1)
        decimal = float(f"{left}.{right}")
        thousands = float(left + right) if len(right) == 3 else decimal
        if len(right) == 3 and reference and reference > 0:
            decimal_distance = abs(math.log(max(decimal, 1e-12) / reference))
            thousands_distance = abs(math.log(max(thousands, 1e-12) / reference))
            value = decimal if decimal_distance <= thousands_distance else thousands
        elif len(right) == 3 and len(left) >= 2:
            value = thousands
        else:
            value = decimal
    else:
        value = float(clean)
    return value * 1000 if kilo else value


def _specs(text: Any, reference: float | None = None, *, default_unit: str = "USD") -> list[PriceSpec]:
    raw_text = _text(text)
    if not raw_text or raw_text in EMPTY:
        return []
    cell_unit = "EUR" if "€" in raw_text and "$" not in raw_text else default_unit
    out: list[PriceSpec] = []
    for match in _PRICE.finditer(raw_text):
        start, end = match.span()
        before = raw_text[start - 1:start]
        after = raw_text[end:end + 1]
        # Dates and percentages are not market prices.
        if before == "/" or after == "/" or after == "%":
            continue
        unit = {"€": "EUR", "$": "USD"}.get(match.group("unit"), cell_unit)
        kilo_low, kilo_high = bool(match.group("ak")), bool(match.group("bk"))
        if match.group("b") and kilo_low != kilo_high:
            # « 89,6–91,2k » : le k ferme la fourchette, il n'appartient pas à
            # une seule borne. Sans cela la borne basse valait 89,6 $ au lieu
            # de 89 600 $ — mille fois trop petite, à côté d'un haut correct.
            kilo_low = kilo_high = True
        first = _number(match.group("a"), reference, kilo_low)
        second = (_number(match.group("b"), reference, kilo_high)
                  if match.group("b") else None)
        low, high = (min(first, second), max(first, second)) if second is not None else (first, None)
        op = "ABOVE" if match.group("op") == ">" else "AT"
        out.append(PriceSpec(
            value=low,
            value_high=high,
            unit=unit,
            operator=op,
            approximate=match.group("op") == "~",
            raw=match.group(0).strip(),
        ))
    return out


def _amounts(text: Any) -> list[AmountSpec]:
    raw = _text(text)
    if not raw or raw in EMPTY:
        return []
    found = []
    for match in _AMOUNT.finditer(raw):
        kind = {"€": "EURO", "$": "USD", "%": "PERCENT"}[match.group("u")]
        found.append(AmountSpec(kind, _number(match.group("n")), raw))
    return found or [AmountSpec("NONE", None, raw)]


def _condition(text: str, operator: str = "ABOVE") -> ConditionDraft | None:
    lowered = text.lower()
    timeframe = ("1W" if any(word in lowered for word in ("hebdo", "semaine")) else
                 "1D" if any(word in lowered for word in ("jour", "journali")) else
                 "4H" if "4h" in lowered else "1H" if "1h" in lowered else None)
    if "clôture" in lowered and timeframe:
        kind = "CLOSE"
    elif any(word in lowered for word in ("retest", "pullback")):
        kind = "RETEST"
    elif any(word in lowered for word in ("rebond", "tenir", "tenue", "doit tenir")):
        kind = "HOLD"
    elif any(word in lowered for word in ("confirm", "cassure", "reprise", ">")):
        kind = "OTHER"
    else:
        return None
    return ConditionDraft(kind, timeframe, operator, text)


def _conditional(text: str) -> bool:
    lowered = text.lower()
    return any(word in lowered for word in (
        "si ", "confirm", "retest", "rebond", "cassure", "reprise", "tenir", "tenue", ">",
    ))


def _headers(sheet) -> dict[str, int]:
    return {_text(cell.value): cell.column for cell in sheet[1] if cell.value is not None}


def _row_map(sheet, row: int, headers: dict[str, int]) -> dict[str, Any]:
    return {name: sheet.cell(row, column).value for name, column in headers.items()}


def _raw_hash(value: dict[str, Any]) -> str:
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode()
    return hashlib.sha256(encoded).hexdigest()[:16]


class _Parser:
    def __init__(self, path: Path, fingerprint: str) -> None:
        self.path = path
        self.fingerprint = fingerprint
        self.ambiguities: list[dict[str, str]] = []
        self.checks: list[dict[str, Any]] = []
        self.references: dict[str, float] = {}

    def ambiguity(self, cell: str, value: Any, reason: str) -> None:
        self.ambiguities.append({"cell": cell, "value": _text(value), "reason": reason})

    def add_level(self, levels: list[LevelDraft], spec: PriceSpec, *, kind: str, label: str,
                  source_text: str, action: ActionDraft | None = None,
                  conditions: list[ConditionDraft] | None = None,
                  extra: dict[str, Any] | None = None) -> None:
        key = (kind, round(spec.value, 10), round(spec.value_high or 0, 10), spec.unit)
        if any((row.kind, round(row.value, 10), round(row.value_high or 0, 10), row.unit) == key
               for row in levels):
            # Keep the richest semantics on the existing level.
            existing = next(row for row in levels if (
                row.kind, round(row.value, 10), round(row.value_high or 0, 10), row.unit) == key)
            existing.action = existing.action or action
            for condition in conditions or []:
                if condition not in existing.conditions:
                    existing.conditions.append(condition)
            return
        metadata = {
            "operator": spec.operator,
            "approximate": spec.approximate,
            "raw_price": spec.raw,
            **(extra or {}),
        }
        levels.append(LevelDraft(kind, spec.value, spec.value_high, spec.unit, label,
                                 source_text, conditions=conditions or [], action=action,
                                 extra=metadata))

    def context_levels(self, text: str, asset: str, levels: list[LevelDraft], cell: str) -> None:
        reference = self.references.get(asset)
        for clause in re.split(r"\s*;\s*", text):
            lowered = clause.lower()
            if not any(word in lowered for word in (
                "support", "résistance", "resistance", "pivot", "zone", "cassure", "trigger",
                "boss final", "bas de range",
            )):
                continue
            if "boss final" in lowered:
                kind, label = "TARGET", "Objectif final"
            elif "résistance" in lowered or "resistance" in lowered:
                kind, label = "RESISTANCE", "Résistance"
            elif "cassure" in lowered or "trigger" in lowered:
                kind, label = "CONFIRMATION", "Cassure / confirmation"
            elif "zone" in lowered or "pivot" in lowered:
                kind, label = "WATCH", "Zone pivot / intérêt"
            else:
                kind, label = "SUPPORT", "Support"
            for spec in _specs(clause, reference):
                self.add_level(levels, spec, kind=kind, label=label, source_text=clause,
                               extra={"source_cell": cell, "informational": True})

    def history_buy(self, text: str, asset: str, levels: list[LevelDraft], cell: str) -> None:
        if not text or text in EMPTY or (any(phrase in text.lower() for phrase in (
            "pas d'achat", "pause nouveaux achats", "attendre", "aucun achat",
        )) and not any(ch.isdigit() for ch in text)):
            return
        reference = self.references.get(asset)
        entry_number = sum(level.kind in {"BUY_ZONE", "REINFORCEMENT"} for level in levels)
        for clause in re.split(r"\s*;\s*", text):
            lowered = clause.lower()
            if not any(ch.isdigit() for ch in clause):
                continue
            amount: AmountSpec | None = None
            price_part = clause
            if "@" in clause:
                amount_part, price_part = clause.split("@", 1)
                amount = (_amounts(amount_part) or [None])[0]
            else:
                amount_match = _AMOUNT.search(clause)
                if amount_match and any(marker in clause[amount_match.end():] for marker in (">", "~")):
                    amount = _amounts(amount_match.group(0))[0]
                    price_part = clause[amount_match.end():]
                elif amount_match and not any(symbol in clause for symbol in ("@", ">")):
                    # A budget without a price cannot be attached to a level safely.
                    specs_without_amount = _specs(clause[:amount_match.start()] + clause[amount_match.end():],
                                                  reference)
                    if not specs_without_amount:
                        self.ambiguity(cell, clause, "Montant présent sans niveau de prix associé")
                        continue
            specs = _specs(price_part, reference)
            if not specs:
                continue
            for spec in specs:
                conditional = _conditional(clause)
                condition = _condition(clause, spec.operator)
                if spec.operator == "ABOVE" or "confirm" in lowered or "cassure" in lowered:
                    kind, label = "CONFIRMATION", "Confirmation d'achat"
                elif "alerte" in lowered:
                    kind, label = "WATCH", "Alerte"
                else:
                    kind = "BUY_ZONE" if entry_number == 0 else "REINFORCEMENT"
                    label = "Achat" if entry_number == 0 else "Recharge"
                    entry_number += 1
                action = None
                if amount:
                    action = ActionDraft(
                        "BUY", amount.amount_type, amount.amount, amount.raw,
                        clause if conditional else "", not conditional,
                        {"source": "EXCEL_HISTORY"},
                    )
                self.add_level(levels, spec, kind=kind, label=label, source_text=clause,
                               action=action, conditions=[condition] if condition else [],
                               extra={"source_cell": cell})

    def targets(self, text: str, asset: str, levels: list[LevelDraft], cell: str,
                parts: list[str] | None = None, historical: bool = False) -> None:
        if not text or text in EMPTY or "pas de vente" in text.lower():
            return
        reference = self.references.get(asset)
        support_mode = "support" in text.lower()
        specs = _specs(text, reference)
        part_specs = [_amounts(value) for value in (parts or [])]
        for index, spec in enumerate(specs):
            if support_mode:
                kind, label, action = "SUPPORT", "Support", None
            else:
                kind, label = "TARGET", f"TP{sum(level.kind == 'TARGET' for level in levels) + 1}"
                amount = part_specs[index][0] if index < len(part_specs) and part_specs[index] else None
                action = ActionDraft(
                    "TAKE_PROFIT",
                    amount.amount_type if amount else "NONE",
                    amount.amount if amount else None,
                    amount.raw if amount else "",
                    "",
                    not historical,
                    {"historical": historical, "source": "EXCEL"},
                )
            self.add_level(levels, spec, kind=kind, label=label, source_text=text,
                           action=action, extra={"source_cell": cell, "historical": historical})

    def history(self, sheet, matrix: dict[tuple[str, str], tuple[str, str]]) -> list[AnalysisDraft]:
        headers = _headers(sheet)
        drafts: list[AnalysisDraft] = []
        # Establish a reference magnitude before parsing comma-thousands values.
        for row in range(2, sheet.max_row + 1):
            asset = _text(sheet.cell(row, headers.get("Crypto", 2)).value).upper()
            raw_price = sheet.cell(row, headers.get("Prix au moment du traitement (approx.)", 3)).value
            specs = _specs(raw_price)
            if asset and specs:
                self.references[asset] = specs[0].value
        for row in range(2, sheet.max_row + 1):
            values = _row_map(sheet, row, headers)
            asset = _text(values.get("Crypto")).upper()
            if not asset:
                continue
            published = _date(values.get("Date"), cell=f"Historique!A{row}",
                              ambiguities=self.ambiguities)
            if published is None:
                continue
            date_key = published.strftime("%d/%m")
            raw_price = _text(values.get("Prix au moment du traitement (approx.)"))
            price_specs = _specs(raw_price, self.references.get(asset))
            price = price_specs[0].value if price_specs else None
            if price:
                self.references[asset] = price
            context = _text(values.get("Contexte / niveau clé"))
            buy = _text(values.get("Achat / recharge"))
            sell = _text(values.get("Vente / objectifs"))
            source_status = _text(values.get("Statut / source"))
            levels: list[LevelDraft] = []
            self.context_levels(context, asset, levels, f"Historique!D{row}")
            self.history_buy(buy, asset, levels, f"Historique!E{row}")
            self.targets(sell, asset, levels, f"Historique!F{row}", historical=True)
            matrix_cell, matrix_text = matrix.get((asset, date_key), ("", ""))
            if not matrix_text:
                self.ambiguity(f"Historique!A{row}:G{row}", f"{asset} {date_key}",
                               "Analyse historique absente de Matrice_dates")
            raw = {key: _text(value) for key, value in values.items()}
            source_ref = (f"excel:{self.path.name}:Historique!{row}:"
                          f"{_raw_hash({'row': raw, 'matrix': matrix_text})}")
            drafts.append(AnalysisDraft(
                asset=asset,
                published_at=published,
                title=f"Historique Excel · {asset} · {published:%d/%m/%Y}",
                source_ref=source_ref,
                source_type="EXCEL_HISTORY",
                source_status=source_status,
                price_at_video=price,
                stance="WAIT" if "attendre" in (buy + context).lower() else "UNSPECIFIED",
                summary=" · ".join(part for part in (context, buy, sell) if part and part not in EMPTY),
                market_context=context,
                requires_revalidation=True,
                levels=levels,
                extra={
                    "import_version": IMPORT_VERSION,
                    "workbook_sha256": self.fingerprint,
                    "source_sheet": "Historique",
                    "source_row": row,
                    "source_cells": raw,
                    "matrix_cell": matrix_cell,
                    "matrix_text": matrix_text,
                    "date_precision": "DAY",
                },
            ))
        return drafts

    def active(self, sheet, dashboard: dict[str, dict[str, str]]) -> list[AnalysisDraft]:
        headers = _headers(sheet)
        drafts: list[AnalysisDraft] = []
        for row in range(2, sheet.max_row + 1):
            values = _row_map(sheet, row, headers)
            asset = _text(values.get("Crypto")).upper()
            if not asset:
                continue
            published = _date(values.get("Dernière MAJ"), cell=f"Ordres_actifs!U{row}",
                              ambiguities=self.ambiguities)
            if published is None:
                continue
            status = _text(values.get("Statut"))
            notes = _text(values.get("Notes"))
            global_condition = _text(values.get("Confirmation / condition"))
            support = _text(values.get("Invalidation / support"))
            requires_revalidation = any(word in status for word in (
                "REVALIDER", "HISTORIQUE", "SORTIE HISTORIQUE",
            ))
            observation_only = status in {
                "ZONE D'INTÉRÊT", "ATTENDRE / PAS D'ORDRE AVEUGLE", "AUCUN SIGNAL D'ACHAT",
            }
            levels: list[LevelDraft] = []
            entry_number = 0
            for number in range(1, 5):
                price_text = _text(values.get(f"Achat {number}"))
                amount_text = _text(values.get(f"Montant {number}"))
                if not price_text or price_text in EMPTY:
                    if amount_text and amount_text not in EMPTY:
                        amount_column = get_column_letter(headers[f"Montant {number}"])
                        self.ambiguity(f"Ordres_actifs!{amount_column}{row}",
                                       amount_text, "Montant sans niveau d'achat")
                    continue
                specs = _specs(price_text, self.references.get(asset))
                if not specs:
                    price_column = get_column_letter(headers[f"Achat {number}"])
                    self.ambiguity(f"Ordres_actifs!{price_column}{row}", price_text,
                                   "Niveau d'achat illisible")
                    continue
                amounts = _amounts(amount_text)
                for index, spec in enumerate(specs):
                    amount = amounts[index] if index < len(amounts) else (amounts[0] if amounts else None)
                    combined = " ; ".join(part for part in (price_text, amount_text, global_condition)
                                          if part)
                    conditional = _conditional(price_text + " " + amount_text)
                    condition = _condition(combined, spec.operator)
                    alert = amount is not None and amount.amount_type == "NONE" and any(
                        word in amount.raw.lower() for word in ("alerte", "réévaluer", "zone profonde"))
                    if observation_only or alert:
                        kind, label = "WATCH", "Zone d'intérêt / surveillance"
                    elif spec.operator == "ABOVE" or "confirm" in price_text.lower():
                        kind, label = "CONFIRMATION", "Confirmation d'achat"
                    else:
                        kind = "BUY_ZONE" if entry_number == 0 else "REINFORCEMENT"
                        label = "Achat principal" if entry_number == 0 else f"Recharge {entry_number}"
                        entry_number += 1
                    enabled = not (conditional or requires_revalidation or observation_only or alert)
                    action = ActionDraft(
                        "WATCH" if observation_only or alert else "BUY",
                        amount.amount_type if amount else "NONE",
                        amount.amount if amount else None,
                        amount.raw if amount else amount_text,
                        combined if not enabled else "",
                        enabled,
                        {"source": "EXCEL_CURRENT", "buy_column": number},
                    )
                    self.add_level(levels, spec, kind=kind, label=label, source_text=price_text,
                                   action=action, conditions=[condition] if condition else [],
                                   extra={"source_cell": (
                                       f"Ordres_actifs!{get_column_letter(headers[f'Achat {number}'])}{row}")})

            target_texts, part_texts = [], []
            for number in range(1, 5):
                target_text = _text(values.get(f"TP {number}"))
                if target_text and target_text not in EMPTY:
                    target_texts.append(target_text)
                    part_texts.append(_text(values.get(f"Part {number}")))
            for target_text, part_text in zip(target_texts, part_texts, strict=True):
                target_number = target_texts.index(target_text) + 1
                target_column = get_column_letter(headers[f"TP {target_number}"])
                self.targets(target_text, asset, levels, f"Ordres_actifs!{target_column}{row}",
                             [part_text], historical=requires_revalidation)

            self.context_levels(global_condition, asset, levels, f"Ordres_actifs!K{row}")
            self.context_levels(support, asset, levels, f"Ordres_actifs!T{row}")
            self._support_semantics(support, asset, levels, f"Ordres_actifs!T{row}")
            self._note_targets(notes, asset, levels, f"Ordres_actifs!V{row}", requires_revalidation)

            dashboard_row = dashboard.get(asset, {})
            if dashboard_row and dashboard_row.get("status") != status:
                self.checks.append({"kind": "DASHBOARD_STATUS_MISMATCH", "asset": asset,
                                    "dashboard": dashboard_row.get("status"), "orders": status})
            raw = {key: _text(value) for key, value in values.items()}
            source_ref = (f"excel:{self.path.name}:Ordres_actifs!{row}:"
                          f"{_raw_hash({'row': raw, 'dashboard': dashboard_row})}")
            stance = ("NEUTRAL" if "CONSERVER" in status else "WAIT" if any(
                word in status for word in ("ATTENDRE", "REVALIDER", "ZONE", "HISTORIQUE", "AUCUN"))
                else "BUY" if "ORDRES" in status else "UNSPECIFIED")
            drafts.append(AnalysisDraft(
                asset=asset,
                published_at=published,
                title=f"Plan courant Excel · {asset} · {published:%d/%m/%Y}",
                source_ref=source_ref,
                source_type="EXCEL_CURRENT",
                source_status=status,
                price_at_video=self.references.get(asset),
                stance=stance,
                summary=" · ".join(part for part in (status, notes) if part),
                market_context=" · ".join(part for part in (global_condition, support, notes) if part),
                requires_revalidation=requires_revalidation,
                levels=levels,
                extra={
                    "import_version": IMPORT_VERSION,
                    "workbook_sha256": self.fingerprint,
                    "source_sheet": "Ordres_actifs",
                    "source_row": row,
                    "source_cells": raw,
                    "dashboard": dashboard_row,
                    "date_precision": "DAY",
                },
                completed="SORTIE HISTORIQUE" in status,
            ))
        return drafts

    def _support_semantics(self, text: str, asset: str, levels: list[LevelDraft], cell: str) -> None:
        """Cases whose meaning is more specific than the generic keyword pass."""

        lowered = text.lower()
        specs = _specs(text, self.references.get(asset))
        if not specs:
            return
        if "si clôture <" in lowered:
            first, *rest = specs
            cond = ConditionDraft("OTHER", None, "BELOW", text)
            self.add_level(levels, first, kind="INVALIDATION", label="Invalidation si clôture dessous",
                           source_text=text, conditions=[cond], extra={"source_cell": cell})
            for spec in rest:
                self.add_level(levels, spec, kind="SUPPORT", label="Support inférieur",
                               source_text=text, extra={"source_cell": cell})
        elif "boss final" in lowered:
            for spec in specs:
                self.add_level(levels, spec, kind="TARGET", label="Objectif final",
                               source_text=text, extra={"source_cell": cell})

    def _note_targets(self, text: str, asset: str, levels: list[LevelDraft], cell: str,
                      historical: bool) -> None:
        lowered = text.lower()
        marker = next((word for word in ("extensions", "autres tp") if word in lowered), None)
        if marker is None:
            return
        body = text.split(":", 1)[1] if ":" in text else text
        chunks = re.split(r"\s+puis\s+|\s+/\s+", body, flags=re.IGNORECASE)
        for chunk in chunks:
            specs = _specs(chunk, self.references.get(asset))
            percentages = [amount for amount in _amounts(chunk) if amount.amount_type == "PERCENT"]
            for spec in specs:
                pct = percentages[0] if percentages else None
                action = ActionDraft("TAKE_PROFIT", "PERCENT" if pct else "NONE",
                                     pct.amount if pct else None, chunk, "", not historical,
                                     {"source": "EXCEL_NOTE", "historical": historical})
                self.add_level(levels, spec, kind="TARGET",
                               label=f"TP{sum(level.kind == 'TARGET' for level in levels) + 1}",
                               source_text=text, action=action,
                               extra={"source_cell": cell, "extension": True})


def parse_workbook(path: str | Path) -> WorkbookDraft:
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    fingerprint = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, data_only=False, read_only=False)
    missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Feuilles obligatoires absentes : {', '.join(missing)}")
    parser = _Parser(source, fingerprint)

    # Seed each asset with an unambiguous price from the active table before
    # reading French comma-thousands such as ETH ``2,667`` versus XRP ``1,505``.
    orders_sheet = workbook["Ordres_actifs"]
    order_headers = _headers(orders_sheet)
    for row in range(2, orders_sheet.max_row + 1):
        asset = _text(orders_sheet.cell(row, order_headers.get("Crypto", 1)).value).upper()
        if not asset:
            continue
        for number in range(1, 5):
            raw = _text(orders_sheet.cell(row, order_headers.get(f"Achat {number}", 1)).value)
            if not raw:
                continue
            specs = _specs(raw)
            if specs:
                parser.references[asset] = specs[0].value
                break

    dashboard_sheet = workbook["Dashboard"]
    dashboard: dict[str, dict[str, str]] = {}
    for row in range(4, dashboard_sheet.max_row + 1):
        asset = _text(dashboard_sheet.cell(row, 4).value).upper()
        if asset:
            dashboard[asset] = {
                "cell": f"Dashboard!D{row}:F{row}",
                "status": _text(dashboard_sheet.cell(row, 5).value),
                "comment": _text(dashboard_sheet.cell(row, 6).value),
            }

    matrix_sheet = workbook["Matrice_dates"]
    matrix: dict[tuple[str, str], tuple[str, str]] = {}
    for row in range(2, matrix_sheet.max_row + 1):
        asset = _text(matrix_sheet.cell(row, 1).value).upper()
        if not asset:
            continue
        for column in range(2, matrix_sheet.max_column + 1):
            value = _text(matrix_sheet.cell(row, column).value)
            date_key = _text(matrix_sheet.cell(1, column).value)
            if value:
                matrix[(asset, date_key)] = (f"Matrice_dates!{matrix_sheet.cell(row, column).coordinate}",
                                             value)

    history = parser.history(workbook["Historique"], matrix)
    active = parser.active(workbook["Ordres_actifs"], dashboard)
    history_keys = {(draft.asset, draft.published_at.strftime("%d/%m")) for draft in history}
    for key, (cell, value) in matrix.items():
        if key not in history_keys:
            parser.ambiguity(cell, value, "Cellule de matrice sans ligne Historique correspondante")

    active_assets = {draft.asset for draft in active}
    dashboard_assets = set(dashboard)
    parser.checks.append({
        "kind": "DASHBOARD_VS_ACTIVE_ASSETS",
        "ok": active_assets == dashboard_assets,
        "dashboard_only": sorted(dashboard_assets - active_assets),
        "orders_only": sorted(active_assets - dashboard_assets),
    })
    parser.checks.append({
        "kind": "MATRIX_VS_HISTORY",
        "ok": set(matrix) == history_keys,
        "matrix_cells": len(matrix),
        "history_rows": len(history_keys),
    })
    analyses = history + active
    sheets = {}
    for sheet_name in REQUIRED_SHEETS:
        sheet = workbook[sheet_name]
        nonempty = sum(1 for row in sheet.iter_rows() if any(cell.value is not None for cell in row))
        sheets[sheet_name] = {
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "nonempty_rows": nonempty,
            "hidden": sheet.sheet_state != "visible",
        }
    dates = sorted({draft.published_at.date().isoformat() for draft in analyses})
    return WorkbookDraft(source, fingerprint, analyses, len(history), len(active), sheets,
                         sorted(active_assets | {draft.asset for draft in history}), dates,
                         parser.ambiguities, parser.checks)


def _metrics(draft: WorkbookDraft) -> dict[str, Any]:
    levels = [level for analysis in draft.analyses for level in analysis.levels]
    actions = [level.action for level in levels if level.action is not None]
    return {
        "sheets_read": len(draft.sheets),
        "cryptos_detected": len(draft.assets),
        "historical_analyses": draft.historical_count,
        "current_plan_snapshots": draft.current_count,
        "analyses_total": len(draft.analyses),
        "dates_detected": len(draft.dates),
        "buy_levels": sum(level.kind in {"BUY_ZONE", "REINFORCEMENT"} for level in levels),
        "buy_actions": sum(action.action == "BUY" for action in actions),
        "zones": sum(level.value_high is not None for level in levels),
        "confirmations": sum(level.kind in {"CONFIRMATION", "BREAKOUT"} for level in levels),
        "take_profits": sum(level.kind in {"TARGET", "TAKE_PROFIT"} for level in levels),
        "invalidations": sum(level.kind == "INVALIDATION" for level in levels),
        "supports": sum(level.kind == "SUPPORT" for level in levels),
        "plans_requiring_revalidation": sum(analysis.requires_revalidation
                                             for analysis in draft.analyses),
        "ambiguous_cells": len(draft.ambiguities),
    }


def import_workbook(path: str | Path, *, apply: bool = False) -> dict[str, Any]:
    """Parse all sheets and optionally append immutable rows to the private DB."""

    draft = parse_workbook(path)
    report: dict[str, Any] = {
        "schema_version": IMPORT_VERSION,
        "file": str(draft.path),
        "file_name": draft.path.name,
        "sha256": draft.fingerprint,
        "status": "DRY_RUN" if not apply else "IMPORTED",
        "sheets": draft.sheets,
        "assets": draft.assets,
        "dates": draft.dates,
        "metrics": _metrics(draft),
        "checks": draft.checks,
        "ambiguities": draft.ambiguities,
        "created_analysis_ids": [],
        "created_current_ids": {},
        "skipped_existing_rows": 0,
    }
    if not apply:
        return report

    created: list[int] = []
    current_ids: dict[str, int] = {}
    with lexa_session() as session:
        previous_import = session.execute(select(LexaImportRow).where(
            LexaImportRow.fingerprint == draft.fingerprint)).scalar_one_or_none()
        if previous_import is not None:
            saved = dict(previous_import.report or report)
            saved["status"] = "ALREADY_IMPORTED"
            return saved
        known_refs = set(session.execute(select(LexaVideoRow.source_ref).where(
            LexaVideoRow.source_ref.like(f"excel:{draft.path.name}:%"))).scalars())
        for analysis in draft.analyses:
            if analysis.source_ref in known_refs:
                report["skipped_existing_rows"] += 1
                continue
            video = LexaVideoRow(
                title=analysis.title,
                source_name="Historique crypto (Excel)",
                published_at=analysis.published_at,
                source_ref=analysis.source_ref,
                source_kind="EXCEL",
                status="ANALYSED",
            )
            session.add(video)
            session.flush()
            row = LexaAnalysisRow(
                video_id=video.id,
                asset=analysis.asset,
                published_at=analysis.published_at,
                price_at_video=analysis.price_at_video,
                quote="USD",
                stance=analysis.stance,
                summary=analysis.summary,
                market_context=analysis.market_context,
                source_type=analysis.source_type,
                source_status=analysis.source_status,
                requires_revalidation=analysis.requires_revalidation,
                extra=analysis.extra,
                status_override="COMPLETED" if analysis.completed else None,
                status_reason="Plan classé historique/sorti dans le fichier Excel."
                if analysis.completed else "",
            )
            session.add(row)
            session.flush()
            created.append(row.id)
            if analysis.source_type == "EXCEL_CURRENT":
                current_ids[analysis.asset] = row.id
            for level in analysis.levels:
                level_row = LexaLevelRow(
                    analysis_id=row.id,
                    kind=level.kind,
                    original_value=level.value,
                    original_high=level.value_high,
                    unit=level.unit,
                    source_text=level.source_text,
                    condition="UNKNOWN",
                    confidence=level.confidence,
                    label=level.label,
                    basis="EXPLICIT",
                    extra=level.extra,
                )
                session.add(level_row)
                session.flush()
                for condition in level.conditions:
                    session.add(LexaConditionRow(
                        level_id=level_row.id,
                        condition_type=condition.condition_type,
                        timeframe=condition.timeframe,
                        operator=condition.operator,
                        required_closes=1,
                        description=condition.description,
                        basis="EXPLICIT",
                    ))
                if level.action is not None:
                    action = level.action
                    session.add(LexaActionRow(
                        analysis_id=row.id,
                        level_id=level_row.id,
                        action=action.action,
                        amount_type=action.amount_type,
                        amount=action.amount,
                        origin="EXCEL_PLAN",
                        raw_text=action.raw_text,
                        condition_text=action.condition_text,
                        execution_enabled=action.execution_enabled,
                        extra=action.extra,
                    ))
        report["created_analysis_ids"] = created
        report["created_current_ids"] = current_ids
        report["metrics"]["analyses_created"] = len(created)
        session.add(LexaImportRow(
            fingerprint=draft.fingerprint,
            file_name=draft.path.name,
            status="IMPORTED",
            report=report,
        ))

    # Current snapshots supersede older versions; old rows remain immutable.
    from .service import record_supersession

    record_supersession(list(current_ids.values()))
    return report


def update_import_report(fingerprint: str, report: dict[str, Any]) -> None:
    """Attach market revalidation to the import ledger without changing source analyses."""

    with lexa_session() as session:
        row = session.execute(select(LexaImportRow).where(
            LexaImportRow.fingerprint == fingerprint)).scalar_one_or_none()
        if row is not None:
            row.report = report


def report_json(report: dict[str, Any]) -> str:
    return json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True, default=str) + "\n"


def drafts_as_dict(path: str | Path) -> list[dict[str, Any]]:
    """Test/debug view; never used as a second source of truth."""

    return [asdict(analysis) for analysis in parse_workbook(path).analyses]
